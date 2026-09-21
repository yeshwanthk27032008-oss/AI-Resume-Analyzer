import os
import re

try:
    import pymupdf  # type: ignore
except ImportError:
    try:
        import fitz as pymupdf  # type: ignore
    except ImportError:
        pymupdf = None  # type: ignore

try:
    from docx import Document  # type: ignore
except ImportError:
    try:
        import docx  # type: ignore
        Document = docx.Document  # type: ignore
    except Exception:
        Document = None  # type: ignore


def clean_text(text: str) -> str:
    """
    Clean and normalize extracted resume or job description text.
    Preserves alphanumeric terms, common technical symbols (C++, C#, .NET),
    and clean paragraph structure while stripping stray control characters.
    """
    if not text:
        return ""

    # Replace fancy bullets, unicode quotes and dashes
    replacements = {
        "\u2022": "\n* ",
        "\u2023": "\n* ",
        "\u25e6": "\n* ",
        "\u2043": "\n* ",
        "\u2219": "\n* ",
        "\u201c": '"',
        "\u201d": '"',
        "\u2018": "'",
        "\u2019": "'",
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\t": " ",
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)

    # Remove non-printable control characters except standard whitespace
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)

    # Normalize multiple blank lines and spaces
    lines = [re.sub(r"[ ]+", " ", line).strip() for line in text.splitlines()]
    # Filter out excessive consecutive empty lines
    cleaned_lines = []
    prev_empty = False
    for line in lines:
        if not line:
            if not prev_empty:
                cleaned_lines.append("")
                prev_empty = True
        else:
            cleaned_lines.append(line)
            prev_empty = False

    return "\n".join(cleaned_lines).strip()


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file using PyMuPDF."""
    if pymupdf is None:
        raise ImportError("PyMuPDF package is not installed. Please run: pip install pymupdf")

    text_chunks = []
    try:
        with pymupdf.open(file_path) as doc:
            for page_num in range(len(doc)):
                page = doc[page_num]
                page_text = page.get_text("text")
                if page_text:
                    text_chunks.append(page_text)
    except Exception as e:
        raise ValueError(f"Failed to read PDF file: {str(e)}")

    raw_text = "\n".join(text_chunks)
    return clean_text(raw_text)


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx, including paragraphs and tables."""
    if Document is None:
        raise ImportError("python-docx package is not installed. Please run: pip install python-docx")

    text_chunks = []
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text and paragraph.text.strip():
                text_chunks.append(paragraph.text.strip())

        # Also parse table contents which are common in resumes
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_chunks.append(" | ".join(row_texts))
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {str(e)}")

    raw_text = "\n".join(text_chunks)
    return clean_text(raw_text)


def extract_text_from_resume(file_path: str) -> str:
    """
    Dispatcher function to extract clean text from either PDF or DOCX resume.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Resume file not found at {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        extracted = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        extracted = extract_text_from_docx(file_path)
    elif ext == ".doc":
        raise ValueError("Legacy .doc format is not supported. Please save or export your resume as .docx or .pdf.")
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Only .pdf and .docx are supported.")

    if not extracted or len(extracted.strip()) < 20:
        raise ValueError("The uploaded file contains little or no extractable text. Please ensure it is not a scanned image PDF.")

    return extracted
